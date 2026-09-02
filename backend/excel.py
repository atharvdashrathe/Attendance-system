"""
excel.py — Excel export logic

Structure:
    excel_exports/
        DSA.xlsx          ← one file per subject
            Lecture 1     ← one sheet per lecture
            Lecture 2
        DBMS.xlsx
            Lecture 1
            ...

Each sheet contains:
    Roll No | Name | Date | Time | Status | Manually Edited
"""

import os
import re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter

# Directory where Excel files are saved
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL_DIR = os.path.join(BASE_DIR, "excel_exports")

os.makedirs(EXCEL_DIR, exist_ok=True)

# --------------------------------------------------
# Color palette
# --------------------------------------------------

COLOR_HEADER_BG    = "1A56DB"   # deep blue header
COLOR_HEADER_FG    = "FFFFFF"   # white text
COLOR_PRESENT_BG   = "D1FAE5"   # light green
COLOR_PRESENT_FG   = "065F46"   # dark green
COLOR_ABSENT_BG    = "FEE2E2"   # light red
COLOR_ABSENT_FG    = "991B1B"   # dark red
COLOR_MANUAL_BG    = "FEF3C7"   # light amber — manually edited row
COLOR_ALT_ROW      = "F8FAFC"   # alternate row background

THIN_BORDER = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)


def _safe_sheet_name(lecture):
    """Sanitize lecture name for Excel sheet name (max 31 chars, no special chars)."""
    safe = re.sub(r'[\\/*?:\[\]]', '_', lecture)
    return safe[:31]


def _apply_header(ws, headers, col_widths):
    """Write styled header row."""
    header_fill = PatternFill("solid", fgColor=COLOR_HEADER_BG)
    header_font = Font(bold=True, color=COLOR_HEADER_FG, size=11)
    header_align = Alignment(horizontal="center", vertical="center")

    for col_idx, (header, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 28


def write_session_to_excel(subject_name, lecture, date_str, time_str, records):
    """
    Write (or rewrite) the attendance for one session into:
        excel_exports/{subject_name}.xlsx  →  sheet: {lecture}

    Args:
        subject_name : str   e.g. "DSA"
        lecture      : str   e.g. "Lecture 1"
        date_str     : str   e.g. "2026-08-17"
        time_str     : str   e.g. "18:30:15"
        records      : list of dicts:
            {roll_no, name, status, confidence, manually_edited}
    """
    # Build file path
    safe_subject = re.sub(r'[\\/*?:<>|"]', '_', subject_name)
    file_path = os.path.join(EXCEL_DIR, f"{safe_subject}.xlsx")

    # Load or create workbook
    if os.path.exists(file_path):
        wb = load_workbook(file_path)
    else:
        wb = Workbook()
        # Remove the default empty sheet
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

    sheet_name = _safe_sheet_name(lecture)

    # Remove existing sheet for this lecture (will be rewritten)
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]

    ws = wb.create_sheet(title=sheet_name)

    # --------------------------------------------------
    # Headers
    # --------------------------------------------------
    headers    = ["Roll No", "Name", "Date", "Time", "Status", "Manually Edited"]
    col_widths = [12,         24,     14,     12,     12,        18]

    _apply_header(ws, headers, col_widths)

    # --------------------------------------------------
    # Freeze header row
    # --------------------------------------------------
    ws.freeze_panes = "A2"

    # --------------------------------------------------
    # Data rows
    # --------------------------------------------------
    for row_idx, rec in enumerate(records, start=2):
        is_present  = rec["status"] == "Present"
        is_manual   = bool(rec.get("manually_edited"))
        is_even_row = (row_idx % 2 == 0)

        # Row background
        if is_manual:
            row_bg = COLOR_MANUAL_BG
        elif is_even_row:
            row_bg = COLOR_ALT_ROW
        else:
            row_bg = "FFFFFF"

        row_fill = PatternFill("solid", fgColor=row_bg)

        # Status badge color
        if is_present:
            status_fill = PatternFill("solid", fgColor=COLOR_PRESENT_BG)
            status_font = Font(bold=True, color=COLOR_PRESENT_FG, size=10)
        else:
            status_fill = PatternFill("solid", fgColor=COLOR_ABSENT_BG)
            status_font = Font(bold=True, color=COLOR_ABSENT_FG, size=10)

        confidence_val = rec.get("confidence")
        conf_str = f"{confidence_val:.4f}" if confidence_val is not None else "—"

        row_data = [
            rec["roll_no"],
            rec.get("name", rec["roll_no"]),
            date_str,
            time_str,
            rec["status"],
            "Yes" if is_manual else "No"
        ]

        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")

            if col_idx == 5:  # Status column
                cell.fill = status_fill
                cell.font = status_font
            else:
                cell.fill = row_fill
                cell.font = Font(size=10)

        ws.row_dimensions[row_idx].height = 20

    wb.save(file_path)
    print(f"[excel] Saved: {file_path}  sheet: {sheet_name}")
    return file_path


def get_subject_excel_path(subject_name):
    """Return the path to a subject's Excel file (may not exist yet)."""
    safe_subject = re.sub(r'[\\/*?:<>|"]', '_', subject_name)
    return os.path.join(EXCEL_DIR, f"{safe_subject}.xlsx")


def list_excel_files():
    """Return list of {subject, filename, path} for all existing Excel files."""
    files = []
    for fname in sorted(os.listdir(EXCEL_DIR)):
        if fname.endswith(".xlsx"):
            subject = os.path.splitext(fname)[0]
            files.append({
                "subject": subject,
                "filename": fname,
                "path": os.path.join(EXCEL_DIR, fname)
            })
    return files
